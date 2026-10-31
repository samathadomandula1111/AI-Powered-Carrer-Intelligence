
from flask import Flask, jsonify, request, send_from_directory
from flask_cors import CORS
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from io import BytesIO
import json
import os
import re
import sqlite3
from datetime import datetime, timezone
from zipfile import BadZipFile

from google.auth.exceptions import GoogleAuthError, TransportError
from google.auth.transport.requests import Request
from google.oauth2.id_token import verify_firebase_token
from lxml.etree import XMLSyntaxError
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from werkzeug.exceptions import HTTPException

app = Flask(__name__, static_folder=None)
MAX_RESUME_SIZE = 5 * 1024 * 1024
app.config["MAX_CONTENT_LENGTH"] = MAX_RESUME_SIZE + 128 * 1024
app.config["DATABASE_PATH"] = os.environ.get(
    "CAREERINFINITY_DB_PATH",
    os.path.join(app.instance_path, "career-infinity.sqlite3"),
)
FIREBASE_PROJECT_ID = os.environ.get("FIREBASE_PROJECT_ID", "carrerinfinity-28ab1")
ALLOWED_ORIGINS = [
    origin.strip()
    for origin in os.environ.get(
        "CORS_ORIGINS",
        "http://127.0.0.1:5000,http://localhost:5000",
    ).split(",")
    if origin.strip()
]

CORS(app, origins=ALLOWED_ORIGINS)

SKILL_ALIASES = {
    "Python": ("python", "python 3", "python3", "python 2", "python2"),
    "Java": ("java",),
    "JavaScript": ("javascript", "js"),
    "TypeScript": ("typescript",),
    "SQL": ("sql",),
    "HTML": ("html",),
    "CSS": ("css",),
    "React": ("react",),
    "Node.js": ("node.js", "nodejs"),
    "Flask": ("flask",),
    "Django": ("django",),
    "Git": ("git",),
    "Docker": ("docker",),
    "AWS": ("aws", "amazon web services"),
    "Azure": ("azure",),
    "Google Cloud": ("google cloud", "gcp"),
    "Machine Learning": ("machine learning", "ml"),
    "LLMs": ("llm", "llms", "large language model", "large language models"),
    "RAG": ("rag", "retrieval augmented generation", "retrieval-augmented generation"),
    "LangChain": ("langchain",),
    "AI Agents": ("ai agent", "ai agents", "agentic ai", "agentic systems"),
    "Workflow Automation": ("workflow automation", "n8n", "zapier", "make.com"),
    "APIs": ("api", "apis", "application programming interface", "application programming interfaces"),
    "Testing": ("testing", "unit tests", "pytest", "test automation"),
    "Playwright": ("playwright",),
    "Selenium": ("selenium",),
    "TensorFlow": ("tensorflow",),
    "PyTorch": ("pytorch",),
    "Data Analysis": ("data analysis", "data analytics"),
    "Excel": ("excel", "microsoft excel"),
    "Power BI": ("power bi",),
    "Communication": ("communication",),
    "Leadership": ("leadership",),
}

CAREER_PROFILES = (
    ("Python Developer", ("Python", "SQL", "Git", "Flask", "Django")),
    ("Data Analyst", ("SQL", "Excel", "Power BI", "Python", "Data Analysis")),
    ("Frontend Developer", ("HTML", "CSS", "JavaScript", "React")),
    ("Automation Engineer", ("Python", "Workflow Automation", "APIs", "Testing", "Playwright")),
    ("Agentic AI Engineer", ("Python", "LLMs", "RAG", "AI Agents", "APIs")),
    ("Machine Learning Engineer", ("Python", "Machine Learning", "TensorFlow", "PyTorch", "SQL")),
    ("Cloud Support Associate", ("AWS", "Azure", "Google Cloud", "Docker", "Git")),
)


def extract_resume_text(filename, content):
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    try:
        if extension == "pdf":
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise ValueError("Encrypted PDFs are not supported.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif extension == "docx":
            document = Document(BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            paragraphs.extend(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            )
            text = "\n".join(paragraphs)
        else:
            raise ValueError("Please upload a PDF or DOCX resume.")
    except (BadZipFile, PackageNotFoundError, PdfReadError, XMLSyntaxError) as error:
        raise ValueError("The file is invalid or could not be read.") from error

    text = text.strip()
    if not text:
        raise ValueError("No readable text was found in the resume.")
    return text


def analyze_resume(text, target_role=""):
    normalized_text = text.lower()
    skills = []
    for name, aliases in SKILL_ALIASES.items():
        count = sum(
            len(re.findall(r"(?<!\w)" + re.escape(alias) + r"(?!\w)", normalized_text))
            for alias in aliases
        )
        if count:
            level = "Strong" if count >= 4 else "Intermediate" if count >= 2 else "Beginner"
            skills.append({
                "name": name,
                "level": level,
                "reason": (
                    f"Appears {count} times in your resume."
                    if count > 1 else "Appears once in your resume."
                ),
                "score": min(100, 40 + count * 15),
            })

    skill_names = {skill["name"] for skill in skills}
    career_matches = []
    for title, required_skills in CAREER_PROFILES:
        matched = [skill for skill in required_skills if skill in skill_names]
        gaps = [skill for skill in required_skills if skill not in skill_names]
        career_matches.append({
            "title": title,
            "match": round(len(matched) / len(required_skills) * 100),
            "reason": f"Matched {len(matched)} of {len(required_skills)} core skills in your resume.",
            "matched_skills": matched,
            "required_skills": list(required_skills),
            "gaps": gaps,
        })
    career_matches.sort(key=lambda career: career["match"], reverse=True)

    suggestions = []
    if len(text) < 500:
        suggestions.append("Add more detail about your experience, projects, and achievements.")
    if not re.search(r"\b[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}\b", text):
        suggestions.append("Include a professional email address in your contact information.")
    for section, pattern in (
        ("education", r"\beducation\b"),
        ("experience", r"\b(experience|employment|work history)\b"),
        ("projects", r"\bprojects?\b"),
    ):
        if not re.search(pattern, normalized_text):
            suggestions.append(f"Consider adding an {section} section.")
    if not skills:
        suggestions.append("Include a clearly labeled skills section with your relevant technologies.")
    if not re.search(
        r"\b\d+(?:\.\d+)?\s*(%|percent|users|customers|hours|days|weeks|months|years|x)\b",
        normalized_text,
    ):
        suggestions.append("Add a metric to one project bullet, such as users reached, time saved, or accuracy gained.")
    selected_role = next(
        (career for career in CAREER_PROFILES if career[0].casefold() == target_role.casefold()),
        None,
    )
    if selected_role:
        missing_keywords = [
            skill for skill in selected_role[1]
            if skill not in {item["name"] for item in skills}
        ]
        if missing_keywords:
            suggestions.append(
                f"ATS keyword check for {selected_role[0]}: if accurate, include "
                + ", ".join(missing_keywords)
                + " in your skills or project descriptions."
            )
    if not suggestions:
        suggestions.append("Keep your resume tailored to each role and quantify your achievements.")

    sections = (
        r"\beducation\b",
        r"\b(experience|employment|work history)\b",
        r"\bprojects?\b",
        r"\bskills?\b",
    )
    structure_score = round(sum(bool(re.search(pattern, normalized_text)) for pattern in sections) / len(sections) * 100)
    project_score = min(100, 35 + (35 if re.search(r"\bprojects?\b", normalized_text) else 0)
                        + (30 if len(re.findall(r"\b(project|built|created|developed|implemented)\b", normalized_text)) >= 3 else 0))
    impact_score = 80 if re.search(r"\b\d+(?:\.\d+)?\s*(%|percent|users|customers|hours|days|weeks|months|years|x)\b", normalized_text) else 35
    keyword_score = min(100, 30 + min(len(skills), 7) * 10)
    skills_score = min(100, 25 + min(len(skills), 6) * 12)
    resume_breakdown = {
        "skills": skills_score,
        "projects": project_score,
        "impact": impact_score,
        "structure": structure_score,
        "keywords": keyword_score,
    }
    structure_sections = sum(
        bool(re.search(pattern, normalized_text)) for pattern in sections
    )
    has_project_section = bool(re.search(r"\bprojects?\b", normalized_text))
    has_action_terms = len(re.findall(
        r"\b(project|built|created|developed|implemented)\b", normalized_text
    )) >= 3
    has_impact_metrics = bool(re.search(
        r"\b\d+(?:\.\d+)?\s*(%|percent|users|customers|hours|days|weeks|months|years|x)\b",
        normalized_text,
    ))
    resume_explanations = {
        "skills": f"{len(skills)} skills from our keyword list appear in your resume.",
        "projects": (
            "A projects section and several project/action terms were found."
            if has_project_section and has_action_terms
            else "Add a projects section and describe what you built."
        ),
        "impact": (
            "We found at least one measurable result or scale."
            if has_impact_metrics
            else "Add numbers to show the result or scale of your work."
        ),
        "structure": f"{structure_sections} of {len(sections)} common section headings were found.",
        "keywords": (
            f"{len(skills)} recognized skill keywords were found for {target_role}."
            if target_role else f"{len(skills)} recognized skill keywords were found."
        ),
    }
    # Keyword heuristics cannot certify a perfect resume, so scores are capped below 100.
    resume_score = min(95, round(sum(resume_breakdown.values()) / len(resume_breakdown)))
    profile_strength = min(95, round((structure_score + skills_score + project_score) / 3))
    best_match = career_matches[0]["match"] if career_matches else 0

    return {
        "resume_score": resume_score,
        "resume_breakdown": resume_breakdown,
        "resume_explanations": resume_explanations,
        "career_match_score": best_match,
        "profile_strength": profile_strength,
        "skills": skills,
        "career_matches": career_matches,
        "suggestions": suggestions,
        "disclaimer": "Scores and keyword matches are estimates, not a formal assessment.",
    }


DEFAULT_PROFILE = {
    "name": "",
    "target_role": "",
    "education": "",
    "graduation_year": "",
    "github": "",
    "linkedin": "",
    "progress": {
        "selectedRole": "",
        "progress": {},
        "bookmarks": [],
        "addedSkills": [],
        "completedSkills": [],
    },
    "drives": [],
    "interview_answers": {},
}


def get_database():
    database_path = app.config["DATABASE_PATH"]
    os.makedirs(os.path.dirname(os.path.abspath(database_path)), exist_ok=True)
    connection = sqlite3.connect(database_path, timeout=10)
    connection.row_factory = sqlite3.Row
    connection.execute("""
        CREATE TABLE IF NOT EXISTS user_profiles (
            uid TEXT PRIMARY KEY,
            profile_json TEXT NOT NULL,
            analysis_json TEXT,
            updated_at TEXT NOT NULL
        )
    """)
    return connection


def load_profile(uid, connection):
    row = connection.execute(
        "SELECT profile_json, analysis_json FROM user_profiles WHERE uid = ?",
        (uid,),
    ).fetchone()
    if row is None:
        return dict(DEFAULT_PROFILE), None
    profile = dict(DEFAULT_PROFILE)
    profile.update(json.loads(row["profile_json"]))
    analysis = json.loads(row["analysis_json"]) if row["analysis_json"] else None
    return profile, analysis


def save_profile(uid, profile, analysis, connection):
    connection.execute(
        """
        INSERT INTO user_profiles (uid, profile_json, analysis_json, updated_at)
        VALUES (?, ?, ?, ?)
        ON CONFLICT(uid) DO UPDATE SET
            profile_json = excluded.profile_json,
            analysis_json = excluded.analysis_json,
            updated_at = excluded.updated_at
        """,
        (
            uid,
            json.dumps(profile),
            json.dumps(analysis) if analysis is not None else None,
            datetime.now(timezone.utc).isoformat(),
        ),
    )


def authenticated_uid():
    authorization = request.headers.get("Authorization", "")
    if not authorization.startswith("Bearer ") or not authorization[7:].strip():
        return None, (jsonify({"error": "Please sign in before continuing."}), 401)
    try:
        claims = verify_firebase_token(
            authorization[7:].strip(),
            Request(),
            audience=FIREBASE_PROJECT_ID,
        )
    except TransportError:
        return None, (jsonify({"error": "Authentication verification is temporarily unavailable. Please try again."}), 503)
    except (GoogleAuthError, ValueError):
        return None, (jsonify({"error": "Your sign-in session is invalid or expired. Please sign in again."}), 401)
    if (
        claims.get("iss") != f"https://securetoken.google.com/{FIREBASE_PROJECT_ID}"
        or not claims.get("sub")
    ):
        return None, (jsonify({"error": "Your sign-in session is invalid or expired. Please sign in again."}), 401)
    return claims["sub"], None


def validate_profile_patch(values):
    if not isinstance(values, dict):
        raise ValueError("Profile changes must be a JSON object.")
    allowed_fields = {
        "name", "target_role", "education", "graduation_year",
        "github", "linkedin", "progress", "drives", "interview_answers",
    }
    if set(values) - allowed_fields:
        raise ValueError("The profile contains unsupported fields.")
    sanitized = {}
    for field in allowed_fields - {"progress", "drives", "interview_answers", "graduation_year"}:
        if field in values:
            value = values[field]
            if not isinstance(value, str) or len(value) > 200:
                raise ValueError(f"{field.replace('_', ' ').capitalize()} must be 200 characters or fewer.")
            sanitized[field] = value.strip()
    if "graduation_year" in values:
        year = values["graduation_year"]
        if year not in ("", None) and (not isinstance(year, int) or year < 1950 or year > 2100):
            raise ValueError("Graduation year must be between 1950 and 2100.")
        sanitized["graduation_year"] = year or ""
    if "progress" in values:
        progress = values["progress"]
        if not isinstance(progress, dict) or set(progress) - {
            "selectedRole", "progress", "bookmarks", "addedSkills", "completedSkills"
        }:
            raise ValueError("Roadmap progress has an invalid format.")
        if not isinstance(progress.get("selectedRole", ""), str):
            raise ValueError("The selected roadmap role must be text.")
        if not isinstance(progress.get("progress", {}), dict):
            raise ValueError("Roadmap progress must be a JSON object.")
        if any(
            not isinstance(progress.get(key, []), list)
            or len(progress.get(key, [])) > 100
            or any(not isinstance(item, str) or len(item) > 160 for item in progress.get(key, []))
            for key in ("bookmarks", "addedSkills", "completedSkills")
        ):
            raise ValueError("Roadmap saved items have an invalid format.")
        sanitized["progress"] = progress
    if "drives" in values:
        drives = values["drives"]
        if not isinstance(drives, list) or len(drives) > 100:
            raise ValueError("Saved opportunities must be a list of up to 100 items.")
        if any(
            not isinstance(item, dict)
            or set(item) - {"company", "details"}
            or not isinstance(item.get("company"), str)
            or not isinstance(item.get("details"), str)
            or len(item["company"]) > 100
            or len(item["details"]) > 300
            for item in drives
        ):
            raise ValueError("Each opportunity needs a company and details.")
        sanitized["drives"] = [
            {"company": item["company"].strip(), "details": item["details"].strip()}
            for item in drives
        ]
    if "interview_answers" in values:
        answers = values["interview_answers"]
        if (
            not isinstance(answers, dict)
            or len(answers) > 50
            or any(
                not isinstance(question_id, str)
                or not re.fullmatch(r"[a-z0-9-]{1,100}", question_id)
                or not isinstance(answer, str)
                or len(answer) > 4000
                for question_id, answer in answers.items()
            )
        ):
            raise ValueError("Interview answers must be up to 50 valid answers of 4,000 characters or fewer.")
        sanitized["interview_answers"] = answers
    return sanitized


@app.route("/")
@app.route("/index.html")
def home():
    return send_from_directory('.', 'index.html')


@app.route("/dashboard")
@app.route("/dashboard.html")
def dashboard():
    return send_from_directory('.', 'dashboard.html')

@app.route("/api/health")
def health():
    return jsonify({
        "status": "healthy",
        "project": "CareerInfinity"
    })


@app.errorhandler(HTTPException)
def handle_http_exception(error):
    if request.path.startswith("/api/"):
        return jsonify({"error": error.description}), error.code
    return error


@app.errorhandler(Exception)
def handle_unexpected_error(error):
    app.logger.exception("Unhandled application error")
    if request.path.startswith("/api/"):
        return jsonify({"error": "An unexpected server error occurred."}), 500
    return "An unexpected server error occurred.", 500


@app.errorhandler(413)
def request_too_large(_error):
    return jsonify({"error": "Maximum resume size is 5 MB."}), 413


@app.route("/api/profile", methods=["GET", "PATCH"])
def user_profile():
    uid, auth_error = authenticated_uid()
    if auth_error:
        return auth_error

    connection = get_database()
    try:
        profile, analysis = load_profile(uid, connection)
        if request.method == "PATCH":
            payload = request.get_json(silent=True)
            if not isinstance(payload, dict) or set(payload) != {"profile"}:
                return jsonify({"error": "Send profile changes in a 'profile' object."}), 400
            try:
                updates = validate_profile_patch(payload["profile"])
            except ValueError as error:
                return jsonify({"error": str(error)}), 400
            if "progress" in updates:
                profile["progress"] = {**profile.get("progress", {}), **updates["progress"]}
            profile.update({key: value for key, value in updates.items() if key != "progress"})
            with connection:
                save_profile(uid, profile, analysis, connection)
        if analysis and profile.get("target_role"):
            analysis_data = analysis.get("data", {})
            target_match = next(
                (career["match"] for career in analysis_data.get("career_matches", [])
                 if career["title"].casefold() == profile["target_role"].casefold()),
                0,
            )
            analysis_data["career_match_score"] = target_match
        return jsonify({"profile": profile, "analysis": analysis})
    finally:
        connection.close()


@app.route("/api/analyze-resume", methods=["POST"])
def analyze_resume_upload():
    uid, auth_error = authenticated_uid()
    if auth_error:
        return auth_error

    resume = request.files.get("resume")
    if resume is None or not resume.filename:
        return jsonify({"error": "Choose a resume to upload."}), 400

    content = resume.read(MAX_RESUME_SIZE + 1)
    if len(content) > MAX_RESUME_SIZE:
        return jsonify({"error": "Maximum resume size is 5 MB."}), 413

    try:
        text = extract_resume_text(resume.filename, content)
    except ValueError as error:
        return jsonify({"error": str(error)}), 400

    connection = get_database()
    try:
        profile, _previous_analysis = load_profile(uid, connection)
        target_role = profile.get("target_role", "").strip()
        result = analyze_resume(text, target_role)
        if target_role:
            target_match = next(
                (career["match"] for career in result["career_matches"]
                 if career["title"].casefold() == target_role.casefold()),
                0,
            )
            result["career_match_score"] = target_match
        analysis = {
            "filename": os.path.basename(resume.filename),
            "analyzed_at": datetime.now(timezone.utc).isoformat(),
            "data": result,
        }
        with connection:
            save_profile(uid, profile, analysis, connection)
        return jsonify(result)
    finally:
        connection.close()


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    host = os.environ.get("HOST", "127.0.0.1")
    app.run(host=host, port=port, debug=False)
